"""evalx.mutations -- paired mutation harness (Agent F).

Implements ONE paired mutation end-to-end (PLAN §3.8) on a COPY of the
practice dossier:

    data/mutated_dup_payment/fraud/   -- duplicate-payment injection:
        vendor 200100 (Pollux Verpackung e.K.), invoice ER900026
        (gross 44,536.94 EUR, originally paid once on 07.03.2025 under
        AZ6600449 / GL entry 7701824) is paid a SECOND time on 21.03.2025
        (new GL entry 7708392, payment run AZ6602901, same BELEGNUMMER
        ER900026). Propagated to: GL (2 rows), Lieferantenbuchungen.txt
        (1 row), OP-Liste_Kreditoren_2025.xlsx (Saldo 200100
        -55,540.87 -> -11,003.93).

    data/mutated_dup_payment/repair/  -- innocent-repair twin:
        the SAME duplicate payment, plus a vendor refund backed by credit
        note GS900026 on 28.03.2025 (GL entry 7708393: debit bank 271000,
        credit 330000-200100; 1 offsetting Lieferantenbuchungen row).
        Net effect zero -> OP list unchanged. A symmetric defense pass must
        flip its verdict on this variant (innocence predicate: offsetting
        credit note / refund exists for the second payment).

Invariants checked by `verify()` in BOTH variants (and the baseline):
    1. GL total balance == 0.00 (Decimal, zero tolerance)
    2. every ENTRY_ID group (semantic alias, col pos 19) balances to 0.00
    3. injected entry_ids are new and each has exactly its expected rows
    4. sum(vendor subledger) == sum(GL 330000-*)          (global tie)
    5. per-vendor tie for 200100                          (delta consistency)
    6. OP-Liste Saldo 200100 == vendor subledger sum 200100
    7. GL / subledger row-count deltas match the injection plan
    8. every injected document reference resolves (ER900026 exists as a
       posted vendor invoice; GS900026 introduced together with its rows)
    9. injected posting dates inside FY2025 and before the
       Festschreibung lock date 2026-01-20
   10. repair variant: injected rows net to exactly 0.00 in GL vendor
       lines and in the subledger

KNOWN residual inconsistencies (documented, intentional -- contract §6 asks
for GL + vendor_tx + OP propagation):
    * Exportprotokoll_GDPdU_2025.pdf row counts / SHA-256 for the two txt
      files no longer match (ingest's manifest check will flag the mutated
      dossier; run ingest with a manifest override or accept the flag).
    * Saldenliste_2025.xlsx / JA-Entwurf are NOT regenerated: in the fraud
      variant bank 271000 and AP control 330000 GL sums shift by 44,536.94,
      creating a trial-balance<->GL deviation (a B-class recon signal on
      top of the intended duplicate-payment signal).

What a FULL mutation suite would add (documented, not implemented):
    * paired variants for every scheme in scope: fictitious vendor
      (+/- contract & goods receipt & independent approval), threshold
      splitting (+/- batch payment authorization), repair capitalization
      (+/- technical assessment), cutoff (+/- per-invoice accrual),
      bank-change-then-payment redirection, fictitious sales -- including
      types ABSENT from the practice set to measure generalization
    * regeneration of Exportprotokoll row counts/hashes and Saldenliste so
      mutated dossiers are fully self-consistent
    * defense-symmetry assertions run through the real pipeline: the
      repair twin's finding must flip to rejected/observation while the
      fraud twin stays reported
    * randomized placement (vendor / amount / date sampled per seed) to
      prevent the finder from overfitting fixed injection coordinates

Usage:
    python3 -m evalx.mutations [--data data/practice]
                               [--out data/mutated_dup_payment]
                               [--verify-only]
Exit code 0 iff all invariants hold in both variants.
No LLM is used (llm_used: false).
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

from evalx.mutation_suite import dossier_hash

ENC = "cp1252"
GL_REL = "Sachkonten/Sachkontobuchungen.txt"
VT_REL = "Kreditoren/Lieferantenbuchungen.txt"
OP_REL = "Begleitdokumente/OP-Liste_Kreditoren_2025.xlsx"

VENDOR = "200100"
INVOICE = "ER900026"
CREDIT_NOTE = "GS900026"
AMOUNT = Decimal("44536.94")
DUP_ENTRY = "7708392"
REPAIR_ENTRY = "7708393"
DUP_DATE = "21.03.2025"
REPAIR_DATE = "28.03.2025"
LOCK_DATE = date(2026, 1, 20)
BASE_OP_SALDO = Decimal("-55540.87")

BASE_GL_ROWS = 20258
BASE_VT_ROWS = 2584


def _amt(d: Decimal) -> str:
    """German decimal-comma format as used in the txt files (no 1000-sep)."""
    return f"{d:.2f}".replace(".", ",")


def _parse_amt(s: str) -> Decimal:
    return Decimal(s.replace(".", "").replace(",", "."))


# ------------------------------------------------------- injected raw lines
# Field order & quoting copied 1:1 from observed rows of each file.

def _gl_line(account, haben, amount, text, bdate, bnum, belegnr, entry_id,
             dokument, user, edate, etime, periode):
    return (f'"{account}";"{periode}";"";"Vortrag";"Zahlung";"Nein";'
            f'"{haben}";{_amt(amount)};"EUR";{_amt(amount)};"{text}";'
            f'{bdate};"{bnum}";{bdate};"{belegnr}";"Aktuell";"";"";'
            f'{entry_id};"{dokument}";"{user}";{edate};"{etime}";"Ja"')


def _vt_line(account, bnum, bdate, text, amount):
    return (f'"{account}";"{bnum}";{bdate};"";{bdate};"{text}";'
            f'{_amt(amount)};"EUR";{_amt(amount)};"";"";"Purch";"Yes"')


DUP_GL = [
    _gl_line(f"330000-{VENDOR}", "Nein", AMOUNT, "Zahlungsausgang",
             DUP_DATE, "AZ6602901", INVOICE, DUP_ENTRY, INVOICE,
             "MV-U11", DUP_DATE, "07:42:10", "Periode 3"),
    _gl_line("271000", "Ja", -AMOUNT, "Zahlungsausgang",
             DUP_DATE, "AZ6602901", INVOICE, DUP_ENTRY, INVOICE,
             "MV-U11", DUP_DATE, "07:42:10", "Periode 3"),
]
DUP_VT = [_vt_line(VENDOR, INVOICE, DUP_DATE, "Zahlungsausgang", AMOUNT)]

REPAIR_GL = [
    _gl_line("271000", "Nein", AMOUNT,
             f"Rückerstattung Lieferant {INVOICE} Gutschrift {CREDIT_NOTE}",
             REPAIR_DATE, CREDIT_NOTE, CREDIT_NOTE, REPAIR_ENTRY, INVOICE,
             "MV-U03", REPAIR_DATE, "09:14:22", "Periode 3"),
    _gl_line(f"330000-{VENDOR}", "Ja", -AMOUNT,
             f"Rückerstattung Lieferant {INVOICE} Gutschrift {CREDIT_NOTE}",
             REPAIR_DATE, CREDIT_NOTE, CREDIT_NOTE, REPAIR_ENTRY, INVOICE,
             "MV-U03", REPAIR_DATE, "09:14:22", "Periode 3"),
]
REPAIR_VT = [_vt_line(VENDOR, CREDIT_NOTE, REPAIR_DATE,
                      f"Rückerstattung Lieferant {INVOICE}", -AMOUNT)]

VARIANTS = {
    "fraud": {
        "gl_lines": DUP_GL, "vt_lines": DUP_VT,
        "op_saldo": BASE_OP_SALDO + AMOUNT,          # -11003.93
        "entries": {DUP_ENTRY: 2},
        "vendor_delta": AMOUNT,
        "expected_detection": (
            "duplicate payment: BELEGNUMMER ER900026 paid twice "
            "(GL entries 7701824 and 7708392, same vendor, same amount)"),
    },
    "repair": {
        "gl_lines": DUP_GL + REPAIR_GL, "vt_lines": DUP_VT + REPAIR_VT,
        "op_saldo": BASE_OP_SALDO,                   # net zero by year end
        "entries": {DUP_ENTRY: 2, REPAIR_ENTRY: 2},
        "vendor_delta": Decimal("0.00"),
        "expected_detection": (
            "same duplicate payment PLUS offsetting refund GS900026 -> "
            "defense pass must flip the verdict (innocence predicate: "
            "offsetting credit note / refund exists)"),
    },
}


# ------------------------------------------------------------------ build

def _append_lines(path: Path, lines):
    raw = path.read_bytes()
    nl = b"\r\n" if b"\r\n" in raw[:8192] or raw.endswith(b"\r\n") else b"\n"
    if not raw.endswith((b"\n",)):
        raw += nl
    for line in lines:
        raw += line.encode(ENC) + nl
    path.write_bytes(raw)


def build_variant(data_dir: Path, out_dir: Path, name: str):
    spec = VARIANTS[name]
    vdir = out_dir / name
    shutil.copytree(data_dir, vdir,
                    ignore=shutil.ignore_patterns(".DS_Store"))
    _append_lines(vdir / GL_REL, spec["gl_lines"])
    _append_lines(vdir / VT_REL, spec["vt_lines"])
    # OP list adjustment (fraud only changes the Saldo; repair nets to zero)
    if spec["op_saldo"] != BASE_OP_SALDO:
        wb = openpyxl.load_workbook(vdir / OP_REL)
        ws = wb.worksheets[0]
        hit = False
        for row in ws.iter_rows(min_row=4):
            if str(row[0].value) == VENDOR:
                row[3].value = float(spec["op_saldo"])
                hit = True
                break
        if not hit:
            raise RuntimeError(f"vendor {VENDOR} not found in OP list")
        wb.save(vdir / OP_REL)
    manifest = {
        "mutation": "duplicate_payment",
        "variant": name,
        "llm_used": False,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "base_dossier_hash": dossier_hash(data_dir),
        "generator": "evalx.mutations",
        "generator_version": "2.0",
        "vendor": VENDOR, "invoice": INVOICE, "amount_eur": float(AMOUNT),
        "injected_gl_entry_ids": sorted(spec["entries"]),
        "injected_rows": {"gl": len(spec["gl_lines"]),
                          "vendor_tx": len(spec["vt_lines"])},
        "op_saldo_200100": float(spec["op_saldo"]),
        "expected_detection": spec["expected_detection"],
        "innocent_predicate": (
            "offsetting credit note / vendor refund exists"
            if name == "repair" else None
        ),
        "affected_files": [GL_REL, VT_REL] + (
            [OP_REL] if spec["op_saldo"] != BASE_OP_SALDO else []
        ),
        "known_residuals": [
            "Exportprotokoll row counts / SHA-256 for the two txt files no "
            "longer match (intentional; see module docstring)",
            "Saldenliste_2025.xlsx not regenerated (fraud variant shifts "
            "GL 271000 / 330000 vs. trial balance)",
        ],
    }
    (vdir / "MUTATION_MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return vdir


# ------------------------------------------------------------------ verify

def _load_gl(d: Path):
    with open(d / GL_REL, encoding=ENC, newline="") as fh:
        return list(csv.reader(fh, delimiter=";"))


def _load_vt(d: Path):
    with open(d / VT_REL, encoding=ENC, newline="") as fh:
        return list(csv.reader(fh, delimiter=";"))


def _op_saldo(d: Path, account: str):
    wb = openpyxl.load_workbook(d / OP_REL, data_only=True)
    for row in wb.worksheets[0].iter_rows(min_row=4, values_only=True):
        if str(row[0]) == account:
            return Decimal(str(row[3]))
    return None


def verify(variant_dir: Path, name: str | None = None):
    """Check all invariants; returns list of {name, passed, detail}."""
    spec = VARIANTS.get(name or "", {"entries": {}, "gl_lines": [],
                                     "vt_lines": [],
                                     "op_saldo": BASE_OP_SALDO,
                                     "vendor_delta": Decimal("0.00")})
    checks = []

    def ck(cname, passed, detail=""):
        checks.append({"name": cname, "passed": bool(passed),
                       "detail": str(detail)})

    gl = _load_gl(variant_dir)
    vt = _load_vt(variant_dir)

    total = sum(_parse_amt(r[7]) for r in gl)
    ck("gl_total_zero", total == 0, f"sum={total}")

    groups = {}
    for r in gl:
        groups[r[18]] = groups.get(r[18], Decimal(0)) + _parse_amt(r[7])
    unbal = {k: str(v) for k, v in groups.items() if v != 0}
    ck("entries_balanced", not unbal,
       f"{len(groups)} entry groups, unbalanced: {unbal or 'none'}")

    for eid, n_rows in spec["entries"].items():
        rows = [r for r in gl if r[18] == eid]
        ck(f"injected_entry_{eid}", len(rows) == n_rows
           and sum(_parse_amt(r[7]) for r in rows) == 0,
           f"rows={len(rows)} (want {n_rows}), sum=0 required")

    gl_vend = sum(_parse_amt(r[7]) for r in gl if r[0].startswith("330000-"))
    vt_all = sum(_parse_amt(r[6]) for r in vt)
    ck("subledger_gl_tie_all", gl_vend == vt_all,
       f"GL 330000-*={gl_vend} vendor_tx={vt_all}")

    gl_t = sum(_parse_amt(r[7]) for r in gl if r[0] == f"330000-{VENDOR}")
    vt_t = sum(_parse_amt(r[6]) for r in vt if r[0] == VENDOR)
    want_t = BASE_OP_SALDO + spec["vendor_delta"]
    ck("subledger_gl_tie_target", gl_t == vt_t == want_t,
       f"GL={gl_t} sub={vt_t} expected={want_t}")

    op = _op_saldo(variant_dir, VENDOR)
    ck("op_matches_subledger_target", op is not None and op == vt_t,
       f"OP={op} subledger={vt_t}")

    ck("row_deltas",
       len(gl) == BASE_GL_ROWS + len(spec["gl_lines"])
       and len(vt) == BASE_VT_ROWS + len(spec["vt_lines"]),
       f"gl={len(gl)} (want {BASE_GL_ROWS + len(spec['gl_lines'])}), "
       f"vt={len(vt)} (want {BASE_VT_ROWS + len(spec['vt_lines'])})")

    inv_rows = [r for r in gl if r[14] == INVOICE and r[4] == "Kreditorenrechnung"]
    ck("doc_ref_resolves", len(inv_rows) >= 1,
       f"{INVOICE} invoice rows in GL: {len(inv_rows)}")

    def _d(s):
        return datetime.strptime(s, "%d.%m.%Y").date()
    inj = [r for r in gl if r[18] in spec["entries"]]
    dates_ok = all(_d(r[11]).year == 2025 and _d(r[21]) <= LOCK_DATE
                   for r in inj)
    ck("dates_ok", dates_ok, f"{len(inj)} injected rows checked")

    if name == "repair":
        inj_sub = sum(_parse_amt(r[6]) for r in vt
                      if r[1] in (CREDIT_NOTE,) or
                      (r[0] == VENDOR and r[2] in (DUP_DATE, REPAIR_DATE)))
        inj_gl = sum(_parse_amt(r[7]) for r in inj
                     if r[0] == f"330000-{VENDOR}")
        ck("repair_nets_zero", inj_sub == 0 and inj_gl == 0,
           f"subledger injected sum={inj_sub}, GL vendor injected sum={inj_gl}")

    return checks


# -------------------------------------------------------------------- main

def main(argv=None):
    ap = argparse.ArgumentParser(prog="evalx.mutations")
    ap.add_argument("--data", default="data/practice")
    ap.add_argument("--out", default="data/mutated_dup_payment")
    ap.add_argument("--verify-only", action="store_true",
                    help="skip building; verify existing variants")
    args = ap.parse_args(argv)
    data_dir, out_dir = Path(args.data), Path(args.out)

    if not args.verify_only:
        if not data_dir.is_dir():
            print(f"[mutations] FATAL: data dir missing: {data_dir}",
                  file=sys.stderr)
            sys.exit(2)
        if out_dir.exists():
            fresh = not any(out_dir.iterdir())
            marker = (out_dir / "fraud" / "MUTATION_MANIFEST.json").exists()
            if not (fresh or marker):
                print(f"[mutations] FATAL: {out_dir} exists and is not a "
                      f"previous mutation output; refusing to delete",
                      file=sys.stderr)
                sys.exit(2)
            shutil.rmtree(out_dir)
        out_dir.mkdir(parents=True)
        for name in VARIANTS:
            build_variant(data_dir, out_dir, name)
            print(f"[mutations] built {out_dir / name}")

    failed = 0
    print(f"\n{'variant':<9} {'check':<28} {'ok':<4} detail")
    print("-" * 78)
    to_verify = [("baseline", data_dir, None)] + \
                [(n, out_dir / n, n) for n in VARIANTS]
    for label, d, name in to_verify:
        if not d.is_dir():
            print(f"{label:<9} {'exists':<28} FAIL directory missing: {d}")
            failed += 1
            continue
        for c in verify(d, name):
            ok = "ok" if c["passed"] else "FAIL"
            failed += not c["passed"]
            print(f"{label:<9} {c['name']:<28} {ok:<4} {c['detail']}")
    print("-" * 78)
    print(f"[mutations] {'ALL INVARIANTS HOLD' if not failed else str(failed) + ' CHECK(S) FAILED'}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
