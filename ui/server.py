"""FinanceAudit UI server (Agent E).

FastAPI backend serving findings + evidence rendering for the three-pane
audit UI in ui/web/.

Data sources (per CONTRACTS.md §5):
  build/findings.json  build/manifest.json  build/audit.duckdb
Fallback (FA_UI_FIXTURES=1 or build artifacts missing):
  ui/fixtures/*.fixture.json + raw files in data/practice/

Run:  uvicorn ui.server:app --port 8642
"""
from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional

from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import JSONResponse
from fastapi.staticfiles import StaticFiles

ROOT = Path(__file__).resolve().parent.parent
BUILD_DIR = Path(os.environ.get("FA_BUILD_DIR", ROOT / "build"))
DATA_DIR = Path(os.environ.get("FA_DATA_DIR", ROOT / "data" / "practice"))
FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures"
FORCE_FIXTURES = os.environ.get("FA_UI_FIXTURES", "") == "1"

# Current workspace (switchable at runtime via POST /api/workspace).
# BUILD_DIR/DATA_DIR above stay the env-derived defaults.
_WS: Dict[str, Path] = {"build": BUILD_DIR, "data": DATA_DIR}


def _bdir() -> Path:
    return _WS["build"]


def _ddir() -> Path:
    return _WS["data"]

CONTRACT_TABLES = [
    "gl", "vendors", "customers", "vendor_tx", "customer_tx", "assets",
    "asset_tx", "goods_receipts", "goods_issues", "sales_invoices",
    "purchase_invoices_2026", "subsequent_payments", "approval_log",
    "masterdata_changes", "permissions", "credit_limits", "shareholders",
    "trial_balance", "op_debitors_accounts", "op_debitors_items",
    "op_creditors_accounts", "reconciliation",
]

app = FastAPI(title="FinanceAudit UI", docs_url=None, redoc_url=None)


@app.middleware("http")
async def _no_cache_static(request, call_next):
    """Dev/demo-friendly: ES modules otherwise stick in the browser cache and
    UI edits appear to have no effect until a hard reload."""
    resp = await call_next(request)
    if not request.url.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-cache"
    return resp

# --------------------------------------------------------------------------
# data access
# --------------------------------------------------------------------------

_cache: Dict[str, Any] = {}


def _load_json(path: Path) -> Any:
    key = str(path)
    mtime = path.stat().st_mtime
    hit = _cache.get(key)
    if hit and hit[0] == mtime:
        return hit[1]
    data = json.loads(path.read_text(encoding="utf-8"))
    _cache[key] = (mtime, data)
    return data


def build_artifacts_present() -> bool:
    return (_bdir() / "findings.json").exists()


def mode() -> str:
    if FORCE_FIXTURES or not build_artifacts_present():
        return "fixtures"
    return "build"


def get_findings() -> List[dict]:
    if mode() == "build":
        data = _load_json(_bdir() / "findings.json")
        # tolerate both bare list and {"findings": [...]}
        return data["findings"] if isinstance(data, dict) and "findings" in data else data
    return _load_json(FIXTURE_DIR / "findings.fixture.json")


def get_manifest() -> dict:
    if mode() == "build" and (_bdir() / "manifest.json").exists():
        return _load_json(_bdir() / "manifest.json")
    return _load_json(FIXTURE_DIR / "manifest.fixture.json")


def _duck():
    """Per-request DuckDB cursor. The base connection is cached, but every
    caller gets its own cursor: FastAPI sync endpoints run in a threadpool and
    a shared connection object races on cursor state (observed as
    ORDER BY "count_star()" binder errors under concurrent requests)."""
    db = str(_bdir() / "audit.duckdb")
    con = _cache.get(f"_duck_con::{db}")
    if con is None:
        import duckdb
        con = duckdb.connect(db, read_only=True)
        _cache[f"_duck_con::{db}"] = con
    return con.cursor()


def lookup_source(source_id: str) -> Optional[dict]:
    """Return normalized source record: file, kind, row_no, page_index, para_no,
    page_label, display_locator, plus fixture file spec if any."""
    if mode() == "fixtures":
        reg = _load_json(FIXTURE_DIR / "source_registry.fixture.json")
        rec = reg["sources"].get(source_id)
        if rec:
            rec = dict(rec)
            rec["file_spec"] = reg["files"].get(rec["file"])
        return rec
    if not (_bdir() / "audit.duckdb").exists():
        return None
    con = _duck()
    rows = con.execute(
        "SELECT source_id, file, kind, row_no, page_index, page_label, para_no, display_locator "
        "FROM source_registry WHERE source_id = ?", [source_id]).fetchall()
    if not rows:
        return None
    r = rows[0]
    return {"source_id": r[0], "file": r[1], "kind": r[2], "row_no": r[3],
            "page_index": r[4], "page_label": r[5], "para_no": r[6],
            "display_locator": r[7], "file_spec": None}


# --------------------------------------------------------------------------
# rendering helpers
# --------------------------------------------------------------------------

def _esc(s: Any) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if s is not None else "")


def _html_table(header: List[str], rows: List[List[Any]], hi_index: int,
                first_line_no: Optional[int] = None) -> str:
    out = ['<table class="src-table"><thead><tr>']
    if first_line_no is not None:
        out.append('<th class="ln">Zeile</th>')
    out += [f"<th>{_esc(h)}</th>" for h in header]
    out.append("</tr></thead><tbody>")
    for i, row in enumerate(rows):
        cls = ' class="hi"' if i == hi_index else ""
        out.append(f"<tr{cls}>")
        if first_line_no is not None:
            out.append(f'<td class="ln">{first_line_no + i + 1}</td>')
        out += [f"<td>{_esc(c)}</td>" for c in row]
        out.append("</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def render_delimited(path: Path, spec: Optional[dict], row_no: int,
                     context: int = 4) -> dict:
    enc = (spec or {}).get("encoding", "cp1252")
    delim = (spec or {}).get("delimiter", ";")
    quote = (spec or {}).get("quotechar", '"')
    has_header = (spec or {}).get("has_header", True)
    columns = (spec or {}).get("columns")
    with open(path, encoding=enc, errors="replace", newline="") as f:
        lines = f.read().splitlines()
    if row_no < 0 or row_no >= len(lines):
        raise HTTPException(404, f"row {row_no} out of range for {path.name}")
    lo = max(1 if has_header else 0, row_no - context)
    hi = min(len(lines), row_no + context + 1)
    window = lines[lo:hi]
    parsed = list(csv.reader(io.StringIO("\n".join(window)),
                             delimiter=delim, quotechar=quote))
    if has_header:
        header = next(csv.reader(io.StringIO(lines[0]), delimiter=delim, quotechar=quote))
    elif columns:
        header = columns
    else:
        header = [f"Spalte {i+1}" for i in range(len(parsed[0]) if parsed else 0)]
    width = len(header)
    parsed = [r[:width] + [""] * (width - len(r)) for r in parsed]
    return {"kind": "table",
            "html": _html_table(header, parsed, row_no - lo, first_line_no=lo)}


def render_xlsx(path: Path, spec: Optional[dict], row_no: int,
                context: int = 3) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = (spec or {}).get("sheet")
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    header_row = (spec or {}).get("header_row", 1)
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = ["" if v is None else str(v) for v in all_rows[header_row - 1]] \
        if 0 < header_row <= len(all_rows) else []
    lo = max(header_row, row_no - 1 - context)   # 0-based index into all_rows
    hi = min(len(all_rows), row_no + context)
    body = [["" if v is None else v for v in r] for r in all_rows[lo:hi]]
    return {"kind": "table",
            "html": _html_table(header, body, row_no - 1 - lo, first_line_no=lo)}


def render_docx(path: Path, para_no: int, context: int = 2) -> dict:
    import docx
    doc = docx.Document(str(path))
    paras = [p.text for p in doc.paragraphs]
    if para_no >= len(paras):
        raise HTTPException(404, f"paragraph {para_no} out of range")
    lo, hi = max(0, para_no - context), min(len(paras), para_no + context + 1)
    parts = ['<div class="src-doc">']
    for i in range(lo, hi):
        cls = "para hi" if i == para_no else "para"
        parts.append(f'<p class="{cls}"><span class="pno">¶{i+1}</span>{_esc(paras[i])}</p>')
    parts.append("</div>")
    return {"kind": "table", "html": "".join(parts)}


def render_pdf(path: Path, page_index: int, quote: Optional[str],
               precision_hint: Optional[str]) -> dict:
    import fitz
    doc = fitz.open(str(path))
    if page_index >= len(doc):
        raise HTTPException(404, f"page {page_index} out of range")
    page = doc[page_index]
    precision = "page_only"
    rects = []
    if quote:
        rects = page.search_for(quote)
        if rects:
            precision = "quote_rect"
        else:  # degrade: successively shorter word prefixes
            words = quote.split()
            for n in (6, 4, 2, 1):
                if len(words) < n:
                    continue
                rects = page.search_for(" ".join(words[:n]))
                if rects:
                    precision = "word_rect"
                    break
    if precision_hint == "page_only":
        rects, precision = [], "page_only"
    for r in rects:
        page.draw_rect(r + (-2, -2, 2, 2), color=(0.75, 0.29, 0.06),
                       fill=(0.98, 0.78, 0.35), fill_opacity=0.30, width=1.4)
    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
    png = base64.b64encode(pix.tobytes("png")).decode("ascii")
    doc.close()
    return {"kind": "page", "png_base64": png, "highlight": bool(rects),
            "locator_precision": precision, "page_index": page_index,
            "page_label": str(page_index + 1)}


def render_from_duckdb(rec: dict) -> Optional[dict]:
    """Build mode: try to show the typed table row from audit.duckdb."""
    try:
        con = _duck()
        table = None
        prof = con.execute(
            "SELECT table_name FROM column_profiles WHERE file = ? LIMIT 1",
            [rec["file"]]).fetchall()
        if prof:
            table = prof[0][0]
        candidates = [table] if table else CONTRACT_TABLES
        for t in candidates:
            try:
                hit = con.execute(
                    f'SELECT row_id FROM "{t}" WHERE source_id = ?',
                    [rec["source_id"]]).fetchall()
            except Exception:
                continue
            if not hit:
                continue
            row_id = hit[0][0]
            cur = con.execute(
                f'SELECT * FROM "{t}" WHERE row_id BETWEEN ? AND ? ORDER BY row_id',
                [row_id - 4, row_id + 4])
            cols = [d[0] for d in cur.description]
            rows = cur.fetchall()
            hi = next((i for i, r in enumerate(rows)
                       if r[cols.index("row_id")] == row_id), 0)
            return {"kind": "table",
                    "html": _html_table(cols, [list(r) for r in rows], hi)}
    except Exception:
        return None
    return None


# --------------------------------------------------------------------------
# endpoints
# --------------------------------------------------------------------------

def _amount_of(f: dict) -> Optional[float]:
    if isinstance(f.get("amount_eur"), (int, float)):
        return f["amount_eur"]
    for c in f.get("claims", []):
        v = c.get("value")
        if c.get("role") == "core" and v and v.get("kind") == "amount":
            return v.get("value")
    return None


@app.get("/api/findings")
def api_findings():
    out = []
    for f in get_findings():
        claims = f.get("claims", [])
        core = [c for c in claims if c.get("role") == "core"]
        out.append({
            "finding_id": f.get("finding_id"),
            "rank": f.get("rank"),
            "title": f.get("title") or f.get("description", "")[:90],
            "scheme": f.get("scheme"),
            "mechanism": f.get("mechanism"),
            "anomaly_type": f.get("anomaly_type"),
            "reporting_status": f.get("reporting_status"),
            "disposition": f.get("disposition"),
            "entity_label": f.get("entity_label") or f.get("entity_key"),
            "amount_eur": _amount_of(f),
            "n_claims": len(claims),
            "core_supported": sum(1 for c in core if c.get("verdict") == "supported"),
            "core_total": len(core),
            "n_missing_evidence": len(f.get("missing_evidence", [])),
            "n_pbc": len(f.get("pbc_requests", [])),
        })
    return out


@app.get("/api/findings/{finding_id}")
def api_finding(finding_id: str):
    for f in get_findings():
        if f.get("finding_id") == finding_id:
            return f
    raise HTTPException(404, f"finding {finding_id} not found")


@app.get("/api/manifest")
def api_manifest():
    man = dict(get_manifest())
    findings = get_findings()
    disp: Dict[str, int] = {}
    for f in findings:
        d = f.get("disposition", "unknown")
        disp[d] = disp.get(d, 0) + 1
    pbc = [{"finding_id": f.get("finding_id"), "title": f.get("title"),
            "disposition": f.get("disposition"),
            "requests": f.get("pbc_requests", [])}
           for f in findings if f.get("pbc_requests")]
    files = man.get("files", [])
    exp = sum(x.get("expected_units", 0) for x in files)
    got = sum(x.get("parsed_units", 0) for x in files)

    # Flagged amount total = sum of report-tier finding amounts.
    flagged_total = sum(
        a for a in (_amount_of(f) for f in findings if f.get("disposition") == "report")
        if isinstance(a, (int, float)))

    # Citation resolvability comes from the eval artefact when present.
    cit = {}
    try:
        if mode() == "build" and (_bdir() / "eval_report.json").exists():
            cit = (_load_json(_bdir() / "eval_report.json") or {}).get("citations") or {}
    except Exception:
        cit = {}

    man["summary"] = {
        "mode": mode(),
        "parse_coverage_pct": round(100.0 * got / exp, 1) if exp else None,
        "n_files": len(files),
        "n_files_complete": sum(1 for x in files
                                if x.get("parse_coverage") == "complete"),
        "dispositions": disp,
        "report_count": disp.get("report", 0),
        "observation_count": disp.get("observation", 0),
        "quarantine_count": disp.get("quarantine", 0),
        "flagged_amount_total": flagged_total,
        "citations": {
            "resolvability_pct": cit.get("resolvability_pct"),
            "total": cit.get("total"),
            "resolved": cit.get("resolved"),
        },
        "pbc_queue": pbc,
    }
    return man


@app.get("/api/status")
def api_status():
    return {
        "ok": True,
        "mode": mode(),
        "build_dir": str(_bdir()),
        "data_dir": str(_ddir()),
        "data_dir_present": _ddir().exists(),
        "artifacts": {
            "findings_json": (_bdir() / "findings.json").exists(),
            "manifest_json": (_bdir() / "manifest.json").exists(),
            "audit_duckdb": (_bdir() / "audit.duckdb").exists(),
        },
        "n_findings": len(get_findings()),
        "llm_used": (lambda: bool((lambda d: d.get("meta", {}).get("llm_used") if isinstance(d, dict) else False)(
            _load_json(_bdir() / "findings.json"))) if (_bdir() / "findings.json").exists() else False)(),
    }


# --------------------------------------------------------------------------
# workspaces (switch between built dossiers / companies)
# --------------------------------------------------------------------------

_WS_ID = re.compile(r"^build[A-Za-z0-9_]*$")


def _company_of(data_dir: Path) -> str:
    key = f"_company::{data_dir}"
    if key in _cache:
        return _cache[key]
    name = data_dir.name
    idx = data_dir / "Sachkonten" / "index.xml"
    try:
        import xml.etree.ElementTree as ET
        el = ET.parse(idx).getroot().find(".//DataSupplier/Name")
        if el is not None and el.text:
            name = el.text.strip()
    except Exception:
        pass
    _cache[key] = name
    return name


def _workspace_info(d: Path) -> Optional[dict]:
    if not (d / "findings.json").exists() or not (d / "audit.duckdb").exists():
        return None
    try:
        man = json.loads((d / "manifest.json").read_text(encoding="utf-8"))             if (d / "manifest.json").exists() else {}
    except Exception:
        man = {}
    dd = (ROOT / man["data_dir"]) if man.get("data_dir") else _ddir()
    try:
        n = len(json.loads((d / "findings.json").read_text(encoding="utf-8"))
                .get("findings", []))
    except Exception:
        n = None
    return {"id": d.name, "build_dir": str(d), "data_dir": str(dd),
            "data_dir_rel": man.get("data_dir"),
            "company": _company_of(dd),
            "dossier": dd.relative_to(ROOT / "data").as_posix()
                       if str(dd).startswith(str(ROOT / "data")) else dd.name,
            "n_findings": n,
            "generated_at": man.get("generated_at"),
            "current": d.resolve() == _bdir().resolve()}


@app.get("/api/workspaces")
def api_workspaces():
    out = []
    for d in sorted(ROOT.glob("build*")):
        if d.is_dir() and _WS_ID.match(d.name):
            info = _workspace_info(d)
            if info:
                out.append(info)
    return out


@app.post("/api/workspace")
def api_workspace_switch(payload: Dict[str, Any]):
    wsid = str(payload.get("id", ""))
    if not _WS_ID.match(wsid):
        raise HTTPException(422, "invalid workspace id")
    d = (ROOT / wsid).resolve()
    if not str(d).startswith(str(ROOT.resolve())) or not d.is_dir():
        raise HTTPException(404, f"workspace {wsid} not found")
    info = _workspace_info(d)
    if info is None:
        raise HTTPException(409, f"{wsid} has no findings.json/audit.duckdb")
    _WS["build"] = d
    _WS["data"] = Path(info["data_dir"])
    return {"ok": True, "workspace": _workspace_info(d)}


# --------------------------------------------------------------------------
# data explorer (browse the structured DuckDB tables)
# --------------------------------------------------------------------------

_BROWSE_EXTRA = ["gl_accounts", "doc_units", "source_registry",
                 "column_profiles", "source_manifest"]


def _browse_allowed() -> List[str]:
    return CONTRACT_TABLES + [t for t in _BROWSE_EXTRA if t not in CONTRACT_TABLES]


@app.get("/api/tables")
def api_tables():
    if mode() != "build":
        raise HTTPException(409, "data explorer requires build mode (run the pipeline)")
    con = _duck()
    present = {r[0] for r in con.execute(
        "SELECT table_name FROM information_schema.tables "
        "WHERE table_schema = 'main'").fetchall()}
    out = []
    for t in _browse_allowed():
        if t not in present:
            continue
        n = con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
        files = [r[0] for r in con.execute(
            "SELECT DISTINCT file FROM column_profiles WHERE table_name = ? "
            "ORDER BY 1", [t]).fetchall()]
        out.append({"name": t, "rows": n, "files": files})
    return out


@app.get("/api/tables/{name}")
def api_table(name: str,
              limit: int = Query(50, ge=1, le=500),
              offset: int = Query(0, ge=0),
              q: Optional[str] = Query(None)):
    if mode() != "build":
        raise HTTPException(409, "data explorer requires build mode (run the pipeline)")
    if name not in _browse_allowed():
        raise HTTPException(404, f"unknown table {name}")
    con = _duck()
    cols = [r[0] for r in con.execute(
        "SELECT column_name FROM information_schema.columns "
        "WHERE table_schema = 'main' AND table_name = ? "
        "ORDER BY ordinal_position", [name]).fetchall()]
    if not cols:
        raise HTTPException(404, f"table {name} has no columns")
    order = '"row_id"' if "row_id" in cols else f'"{cols[0]}"'
    where, params = "", []
    if q:
        blob = "lower(concat_ws(' ', " + ", ".join(
            f'CAST("{c}" AS VARCHAR)' for c in cols) + "))"
        where = f"WHERE {blob} LIKE ?"
        params = [f"%{q.lower()}%"]
    total = con.execute(
        f'SELECT COUNT(*) FROM "{name}" {where}', params).fetchone()[0]
    rows = con.execute(
        f'SELECT * FROM "{name}" {where} ORDER BY {order} LIMIT ? OFFSET ?',
        params + [limit, offset]).fetchall()
    profiles = [
        {"column": r[0], "declared_role": r[1], "semantic_role": r[2],
         "dtype": r[3], "n_unique": r[4], "is_constant": bool(r[5]),
         "conflict": r[6]}
        for r in con.execute(
            "SELECT raw_column, declared_role, semantic_role, dtype, n_unique, "
            "is_constant, conflict FROM column_profiles WHERE table_name = ? "
            "ORDER BY position", [name]).fetchall()]
    return {"name": name, "columns": cols, "total": total,
            "offset": offset, "limit": limit,
            "rows": [["" if v is None else str(v) for v in r] for r in rows],
            "profiles": profiles}


@app.get("/api/source/{source_id}/render")
def api_render(source_id: str,
               quote: Optional[str] = Query(None),
               precision: Optional[str] = Query(None)):
    rec = lookup_source(source_id)
    if rec is None:
        raise HTTPException(404, f"source_id {source_id} not found in registry")
    file_rel = rec["file"]
    path = _ddir() / file_rel
    caption = rec.get("display_locator") or file_rel
    kind = rec.get("kind")
    suffix = path.suffix.lower()
    result: Optional[dict] = None

    if kind in ("cell_row", "table_row", "row") or (
            kind is None and suffix in (".txt", ".csv", ".xlsx")):
        if mode() == "build":
            result = render_from_duckdb(rec)
        if result is None:
            if not path.exists():
                raise HTTPException(404, f"source file missing: {file_rel}")
            if suffix == ".xlsx":
                result = render_xlsx(path, rec.get("file_spec"), int(rec["row_no"]))
            else:
                result = render_delimited(path, rec.get("file_spec"), int(rec["row_no"]))
    elif suffix == ".docx" or (kind == "para" and suffix != ".pdf"):
        if not path.exists():
            raise HTTPException(404, f"source file missing: {file_rel}")
        result = render_docx(path, int(rec.get("para_no") or 0))
    elif suffix == ".pdf" or kind in ("page", "para"):
        if not path.exists():
            raise HTTPException(404, f"source file missing: {file_rel}")
        if quote is None and kind == "para" and mode() == "build":
            # No explicit quote: fall back to the paragraph's own text from
            # doc_units so para-level PDF citations still get a highlight rect.
            try:
                row = _duck().execute(
                    "SELECT text FROM doc_units WHERE source_id = ?",
                    [source_id]).fetchone()
                if row and row[0]:
                    quote = " ".join(str(row[0]).split()[:12])
            except Exception:
                pass
        result = render_pdf(path, int(rec.get("page_index") or 0), quote, precision)
    else:
        raise HTTPException(422, f"cannot render kind={kind} file={file_rel}")

    result["source_id"] = source_id
    result["file"] = file_rel
    result["caption"] = caption
    return JSONResponse(result)


app.mount("/", StaticFiles(directory=str(Path(__file__).resolve().parent / "web"),
                           html=True), name="web")
