"""Parsers for the 19 Begleitdokumente (csv / xlsx / docx / pdf) -> side tables + doc_units."""
from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import docx as docx_lib
import fitz
import openpyxl

from . import common as _common
from .common import (
    ManifestBuilder,
    SourceRegistry,
    detect_decimal_convention,
    excel_amount,
    parse_amount_conv,
    parse_german_amount,
    parse_german_date,
    read_delimited,
    sha256_file,
)

BEGLEIT = "Begleitdokumente"

GDPDU_TXT_FILES = [
    "Sachkonten/Sachkonten.txt",
    "Sachkonten/Sachkontobuchungen.txt",
    "Debitoren/Kunden.txt",
    "Debitoren/Kundenbuchungen.txt",
    "Kreditoren/Lieferanten.txt",
    "Kreditoren/Lieferantenbuchungen.txt",
    "AV/Anlagen.txt",
    "AV/Anlagenbuchungen.txt",
]


# ------------------------------------------------------------- Exportprotokoll

def parse_exportprotokoll(data_dir: Path):
    """Extract declared (rows, sum, sha256) per GDPdU txt from the Exportprotokoll PDF.

    Returns ({rel_path: {declared_rows, declared_sum_raw, declared_sum_dec, declared_sha256}}, errors).
    Hash hex is split across lines in the PDF text layer -> concatenate hex-only lines.
    """
    path = data_dir / BEGLEIT / "Exportprotokoll_GDPdU_2025.pdf"
    entries, errors = {}, []
    if not path.exists():
        return entries, [f"{path.name} not found"]
    try:
        text = fitz.open(path)[0].get_text()
    except Exception as exc:
        return entries, [f"Exportprotokoll unreadable: {exc}"]
    positions = [(text.find(f), f) for f in GDPDU_TXT_FILES if text.find(f) >= 0]
    positions.sort()
    for (pos, f), nxt in zip(positions, positions[1:] + [(len(text), None)]):
        seg = text[pos + len(f): nxt[0]]
        lines = [ln.strip() for ln in seg.splitlines() if ln.strip()]
        try:
            declared_rows = int(lines[0])
            sum_raw = sum_dec = None
            rest = lines[1:]
            if rest and re.match(r"^-?[\d.]+,\d{2}$", rest[0]):
                sum_raw = rest[0]
                sum_dec = str(Decimal(sum_raw.replace(".", "").replace(",", ".")))
                rest = rest[1:]
            elif rest and rest[0] in ("–", "-", "—"):
                rest = rest[1:]
            hexcat = "".join(ln for ln in rest if re.fullmatch(r"[0-9a-f]+", ln))
            sha = hexcat[:64] if len(hexcat) >= 64 else None
            if sha is None:
                errors.append(f"Exportprotokoll: could not assemble sha256 for {f}")
            entries[f] = {
                "declared_rows": declared_rows,
                "declared_sum_raw": sum_raw,
                "declared_sum_dec": sum_dec,
                "declared_sha256": sha,
            }
        except Exception as exc:
            errors.append(f"Exportprotokoll: failed to parse block for {f}: {exc}")
    missing = [f for f in GDPDU_TXT_FILES if f not in entries]
    if missing:
        errors.append(f"Exportprotokoll: no entries found for {missing}")
    return entries, errors


# ------------------------------------------------------------------ CSV files

# target table -> (filename, [(target_col, source_header, type)]) type: text|amount|date|int
CSV_SPECS = {
    "goods_receipts": ("Wareneingangsliste_2025.csv", [
        ("we_no", "WARENEINGANG_NR", "text"), ("we_date", "WARENEINGANG_DATUM", "date"),
        ("invoice_ref", "RECHNUNGSNUMMER", "text"), ("vendor_account", "KREDITOR", "text"),
        ("vendor_name", "KREDITORNAME", "text"), ("amount", "BETRAG_EUR", "amount"),
        ("note", "BEMERKUNG", "text")]),
    "goods_issues": ("Warenausgangsliste_2025.csv", [
        ("wa_no", "WARENAUSGANG_NR", "text"), ("wa_date", "WARENAUSGANG_DATUM", "date"),
        ("invoice_ref", "RECHNUNGSNUMMER", "text"), ("customer_account", "DEBITOR", "text"),
        ("customer_name", "DEBITORNAME", "text"), ("amount", "BETRAG_EUR", "amount"),
        ("location", "LAGERORT", "text")]),
    "sales_invoices": ("Fakturajournal_2025.csv", [
        ("invoice_no", "RECHNUNGSNUMMER", "text"), ("kind", "ART", "text"),
        ("customer_account", "DEBITOR", "text"), ("customer_name", "DEBITORNAME", "text"),
        ("invoice_date", "FAKTURADATUM", "date"), ("service_date", "LEISTUNGSDATUM", "date"),
        ("amount", "BETRAG_EUR", "amount"), ("currency", "WAEHRUNG", "text"),
        ("note", "BEMERKUNG", "text")]),
    "purchase_invoices_2026": ("Fakturajournal_Januar_2026_Kreditoren.csv", [
        ("invoice_no", "RECHNUNGSNUMMER", "text"), ("kind", "ART", "text"),
        ("vendor_account", "KREDITOR", "text"), ("vendor_name", "KREDITORNAME", "text"),
        ("invoice_date", "FAKTURADATUM", "date"), ("service_date", "LEISTUNGSDATUM", "date"),
        ("amount", "BETRAG_EUR", "amount"), ("currency", "WAEHRUNG", "text"),
        ("note", "BEMERKUNG", "text")]),
    "subsequent_payments": ("Buchungen_Folgeperiode_2026.csv", [
        ("posting_date", "BUCHUNGSDATUM", "date"), ("doc_ref", "BELEG", "text"),
        ("customer_account", "DEBITOR", "text"), ("customer_name", "DEBITORNAME", "text"),
        ("amount", "BETRAG_EUR", "amount"), ("currency", "WAEHRUNG", "text"),
        ("offset_account", "GEGENKONTO_HB", "text"), ("text", "BUCHUNGSTEXT", "text")]),
    "approval_log": ("Freigabe-Log_Journale_2025.csv", [
        ("entry_id", "ERFASSUNGSNUMMER", "text"), ("journal_name", "JOURNALNAME", "text"),
        ("n_lines", "ANZAHL_ZEILEN", "int"), ("sum_abs", "SUMME_ABS_EUR", "amount"),
        ("creator", "ERSTELLER", "text"), ("created_date", "ERFASST_AM", "date"),
        ("created_time", "ERFASST_UM", "text"), ("approver", "FREIGEBER", "text"),
        ("approval_date", "FREIGABEDATUM", "date"), ("status", "FREIGABESTATUS", "text")]),
    "masterdata_changes": ("Stammdatenaenderungen_2025.csv", [
        ("change_date", "DATUM", "date"), ("kind", "ART", "text"),
        ("account", "KONTO", "text"), ("name", "NAME", "text"), ("field", "FELD", "text"),
        ("old_value", "WERT_ALT", "text"), ("new_value", "WERT_NEU", "text"),
        ("changed_by", "GEAENDERT_VON", "text"), ("approved_by", "GENEHMIGT_VON", "text"),
        ("approved", "GENEHMIGT", "text")]),
    "credit_limits": ("Kreditlimitliste_Debitoren_2025.csv", [
        ("customer_account", "DEBITOR", "text"), ("customer_name", "DEBITORNAME", "text"),
        ("credit_limit", "KREDITLIMIT_EUR", "amount"),
        ("utilization", "AUSNUTZUNG_31_12_2025_EUR", "amount"),
        ("status", "STATUS", "text"), ("collateral", "BESICHERUNG", "text")]),
    "shareholders": ("Gesellschafterliste_Beteiligungen.csv", [
        ("name", "NAME", "text"), ("share_pct", "ANTEIL_PROZENT", "amount"),
        ("ic_flag", "IC_KENNZEICHEN", "text"), ("relation", "VERHAELTNIS", "text"),
        ("note", "BEMERKUNG", "text")]),
}

CSV_SCOPES = {
    "goods_receipts": "material/logistics goods receipts only — services are NOT covered by this list",
    "goods_issues": "delivery-related goods issues, 1:1 with FY2025 sales invoices; credit notes legitimately have no issue",
    "sales_invoices": "sales invoice journal FY2025 incl. 1 credit note (kind column)",
    "purchase_invoices_2026": "post-period vendor invoices January 2026 as provided (8 invoices)",
    "subsequent_payments": "customer receipts ONLY — not a complete post-period general ledger; no non-receipt inference permitted",
    "approval_log": "journal-level approvals; coverage of all invoices/payments UNPROVEN — absence of a journal here is not evidence of missing approval",
    "masterdata_changes": "master data changes 2025 as provided",
    "credit_limits": "customer credit limits as of 31.12.2025",
    "shareholders": "shareholders and affiliated companies incl. section header lines (see is_section_header)",
}


def _csv_schema(colmap, extra=()):
    """Full ordered column list a CSV table carries once populated (row dict keys).
    Amount columns expand to <tgt>, <tgt>_raw, <tgt>_dec (see _parse_csv_table)."""
    cols = ["row_id", "source_id"]
    for tgt, _src, typ in colmap:
        cols.append(tgt)
        if typ == "amount":
            cols += [tgt + "_raw", tgt + "_dec"]
    cols += list(extra)
    return cols


# ROBUSTNESS (S2/S4): an ABSENT or empty CSV must still create a DuckDB table with its
# FULL declared schema, so rules that SELECT its columns return 0 rows instead of raising
# a Binder Error (column not found) that would crash the Finder. Populated tables get their
# schema from the row dicts as before — this map is only consulted for the empty case.
CSV_TABLE_SCHEMAS = {
    table: _csv_schema(spec[1],
                       extra=(("section", "is_section_header") if table == "shareholders" else ()))
    for table, spec in CSV_SPECS.items()
}


def _parse_csv_table(data_dir, table, spec, registry, manifest):
    fname, colmap = spec
    rel = f"{BEGLEIT}/{fname}"
    path = data_dir / rel
    # ROBUSTNESS (S2 missing-side-table / S4 renamed file): a declared sidecar that is
    # absent must degrade to an empty table + a 'failed' coverage manifest entry, NOT a
    # FileNotFoundError crash. Downstream rules gate on parse_coverage / observability
    # (e.g. R15 goods-receipt obligation, R19 sequence integrity), so an empty table
    # produces no phantom absence claims.
    if not path.exists():
        manifest.add(file=rel, file_hash=None, size_bytes=None, expected_units=None,
                     parsed_units=0, parser_errors=[f"{rel}: file not provided in this dossier"],
                     parse_coverage="failed", population_scope=CSV_SCOPES.get(table, ""),
                     provided=False)
        return [], (rel, [h for _, h, _ in colmap], [])
    file_hash = sha256_file(path)
    lines = read_delimited(path)
    encoding_used = _common.LAST_READ_ENCODING
    header = [h.strip() for h in lines[0][1]]
    errors = []
    if encoding_used != _common.ENCODING:
        errors.append(f"{rel}: decoded as {encoding_used} (cp1252 failed) — non-default encoding")
    idx = {}
    for tgt, src, typ in colmap:
        if src in header:
            idx[tgt] = (header.index(src), typ)
        else:
            errors.append(f"{rel}: declared header {src!r} not found (headers={header})")
    # ROBUSTNESS (S7 decimal variant): infer each amount column's decimal convention from
    # its own values. Practice columns are German -> 'de' (identical output); a dot-decimal
    # finals column is parsed correctly instead of being silently inflated ~100x.
    amount_conv = {}
    for tgt, (i, typ) in idx.items():
        if typ == "amount":
            col_vals = [(f[i] if i < len(f) else "") for _, f, _ in lines[1:]]
            conv = detect_decimal_convention(col_vals)
            amount_conv[tgt] = conv
            if conv != "de":
                errors.append(f"{rel}: column {tgt!r} parsed as dot-decimal (non-German "
                              "convention detected) — flagged for review")
    rows = []
    section = None
    n_nonempty = 0
    for line_no, fields, raw in lines[1:]:
        if not any(x.strip() for x in fields):
            continue  # pure separator line (counted separately)
        n_nonempty += 1
        sid = registry.register(
            file=rel, file_hash=file_hash, kind="cell_row",
            locator=f"{rel}#row:{line_no}", content=raw,
            display_locator=f"{rel}:row {line_no}", row_no=line_no)
        rec = {"row_id": line_no, "source_id": sid}
        for tgt, (i, typ) in idx.items():
            v = fields[i] if i < len(fields) else ""
            try:
                if typ == "date":
                    rec[tgt] = parse_german_date(v)
                elif typ == "amount":
                    val, raw_a, canon = parse_amount_conv(v, amount_conv.get(tgt, "de"))
                    rec[tgt] = val
                    rec[tgt + "_raw"] = raw_a
                    rec[tgt + "_dec"] = canon
                elif typ == "int":
                    rec[tgt] = int(v) if v.strip() else None
                else:
                    rec[tgt] = v.strip() or None
            except Exception as exc:
                errors.append(f"{rel}:row {line_no}:{tgt}: {exc}")
                rec[tgt] = None
        if table == "shareholders":
            only_name = rec.get("name") and not any(
                rec.get(k) for k in ("share_pct", "ic_flag", "relation", "note"))
            if only_name:
                section = rec["name"]
            rec["section"] = section
            rec["is_section_header"] = bool(only_name)
        rows.append(rec)
    manifest.add(file=rel, file_hash=file_hash, size_bytes=path.stat().st_size,
                 expected_units=n_nonempty, parsed_units=len(rows), parser_errors=errors,
                 population_scope=CSV_SCOPES.get(table, ""))
    raw_rows = [f for _, f, _ in lines[1:] if any(x.strip() for x in f)]
    return rows, (rel, header, raw_rows)


# ----------------------------------------------------------------- XLSX files

def _xl_register(registry, rel, file_hash, sheet, row_no, values):
    content = ";".join("" if v is None else str(v) for v in values)
    return registry.register(
        file=rel, file_hash=file_hash, kind="cell_row",
        locator=f"{rel}#sheet:{sheet}#row:{row_no}", content=content,
        display_locator=f"{rel}:{sheet}:row {row_no}", row_no=row_no)


# ROBUSTNESS: empty-table schemas for the xlsx side tables (not in the GDPdU ORDER map),
# so an ABSENT xlsx yields a 0-row DuckDB table WITH its columns rather than a Binder Error
# when a rule SELECTs from it (mirrors CSV_TABLE_SCHEMAS; see write_table in ingest/run.py).
XLSX_TABLE_SCHEMAS = {
    "permissions": ["row_id", "source_id", "user_id", "department", "can_post", "can_approve",
                    "can_pay", "can_create_vendor", "can_periods", "is_admin", "is_mgmt", "note"],
    "op_debitors_accounts": ["row_id", "source_id", "account", "name", "grp",
                             "balance", "balance_raw", "balance_dec"],
    "op_creditors_accounts": ["row_id", "source_id", "account", "name", "grp",
                              "balance", "balance_raw", "balance_dec"],
    "op_debitors_items": ["row_id", "source_id", "account", "name", "doc_ref", "doc_date",
                          "amount", "amount_raw", "amount_dec"],
    "trial_balance": ["row_id", "source_id", "account", "name", "kind", "opening", "opening_dec",
                      "debit", "debit_dec", "credit", "credit_dec", "closing", "closing_dec",
                      "closing_computed"],
    "reconciliation": ["row_id", "source_id", "label", "value", "value_raw", "value_dec", "file"],
}


def _xlsx_absent(manifest, rel, scope):
    """ROBUSTNESS: record a missing xlsx sidecar as a 'failed'-coverage manifest entry (its
    table stays empty-but-schema'd, downstream tie-outs/rules degrade, no phantom absence)."""
    manifest.add(file=rel, file_hash=None, size_bytes=None, expected_units=None,
                 parsed_units=0, parser_errors=[f"{rel}: file not provided in this dossier"],
                 parse_coverage="failed", population_scope=scope + " — NOT PROVIDED",
                 provided=False)


def parse_xlsx_files(data_dir: Path, registry, manifest):
    tables = {"permissions": [], "op_debitors_accounts": [], "op_debitors_items": [],
              "op_creditors_accounts": [], "trial_balance": [], "reconciliation": []}
    profile_feed = []

    # --- Berechtigungsauswertung (header sheet-row 3, 9 users)
    rel = f"{BEGLEIT}/Berechtigungsauswertung_2025.xlsx"
    path = data_dir / rel
    perm_scope = "permission matrix of the 9 posting-capable users per 31.12.2025 (SoD basis)"
    if not path.exists():
        _xlsx_absent(manifest, rel, perm_scope)
    else:
        fh = sha256_file(path)
        ws = openpyxl.load_workbook(path, data_only=True).active
        errors, parsed = [], 0
        raw_rows = []
        for r_no, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            if row[0] is None:
                continue
            sid = _xl_register(registry, rel, fh, ws.title, r_no, row)
            tables["permissions"].append({
                "user_id": str(row[0]), "department": row[1],
                "can_post": row[2] == "X", "can_approve": row[3] == "X",
                "can_pay": row[4] == "X", "can_create_vendor": row[5] == "X",
                "can_periods": row[6] == "X", "is_admin": row[7] == "X",
                "is_mgmt": (str(row[8]).strip().lower() == "ja") if row[8] else False,
                "note": row[9], "row_id": r_no, "source_id": sid})
            raw_rows.append([("" if v is None else str(v)) for v in row])
            parsed += 1
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=parsed, parsed_units=parsed, parser_errors=errors,
                     population_scope=perm_scope)
        profile_feed.append((rel, "permissions",
                             ["Benutzer", "Abteilung", "Buchen", "Journal freigeben", "Zahlungslauf",
                              "Stammdaten/Kreditor anlegen", "Perioden", "Systemadmin", "Management",
                              "Bemerkung"], raw_rows))

    # --- OP-Liste Debitoren (2 sheets) / Kreditoren (1 sheet)
    for rel_name, acc_table, item_table, scope in [
        ("OP-Liste_Debitoren_2025.xlsx", "op_debitors_accounts", "op_debitors_items",
         "AR open items per 31.12.2025; account balances complete, item sheet is an excerpt (Auszug) of current-year items"),
        ("OP-Liste_Kreditoren_2025.xlsx", "op_creditors_accounts", None,
         "AP open account balances per 31.12.2025"),
    ]:
        rel = f"{BEGLEIT}/{rel_name}"
        path = data_dir / rel
        if not path.exists():
            _xlsx_absent(manifest, rel, scope)
            continue
        fh = sha256_file(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        errors, parsed = [], 0
        ws = wb.worksheets[0]
        for r_no, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
            if row[0] is None:
                continue
            sid = _xl_register(registry, rel, fh, ws.title, r_no, row)
            val, raw_a, canon = excel_amount(row[3])
            tables[acc_table].append({
                "account": str(row[0]), "name": row[1], "grp": row[2],
                "balance": val, "balance_raw": raw_a, "balance_dec": canon,
                "row_id": r_no, "source_id": sid})
            parsed += 1
        if item_table and len(wb.worksheets) > 1:
            ws2 = wb.worksheets[1]
            for r_no, row in enumerate(ws2.iter_rows(min_row=4, values_only=True), start=4):
                if row[0] is None:
                    continue
                sid = _xl_register(registry, rel, fh, ws2.title, r_no, row)
                val, raw_a, canon = excel_amount(row[4])
                tables[item_table].append({
                    "account": str(row[0]), "name": row[1], "doc_ref": row[2],
                    "doc_date": parse_german_date(row[3]),
                    "amount": val, "amount_raw": raw_a, "amount_dec": canon,
                    "row_id": r_no, "source_id": sid})
                parsed += 1
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=parsed, parsed_units=parsed, parser_errors=errors,
                     population_scope=scope)

    # --- Saldenliste 2025 (trial balance; closing column is an uncached formula =D+E-F)
    rel = f"{BEGLEIT}/Saldenliste_2025.xlsx"
    path = data_dir / rel
    tb_scope = "trial balance FY2025 (43 accounts) + net-income reconciliation sheet"
    if not path.exists():
        _xlsx_absent(manifest, rel, tb_scope)
    else:
        fh = sha256_file(path)
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Saldenliste 2025"]
        errors, parsed = [], 0
        for r_no, row in enumerate(ws.iter_rows(min_row=5, values_only=True), start=5):
            if row[0] is None:
                continue  # skips the formula-only 'Summe' row
            sid = _xl_register(registry, rel, fh, ws.title, r_no, row)
            op, op_raw, op_dec = excel_amount(row[3])
            de, de_raw, de_dec = excel_amount(row[4])
            cr, cr_raw, cr_dec = excel_amount(row[5])
            closing_dec = str(Decimal(op_dec or "0") + Decimal(de_dec or "0") - Decimal(cr_dec or "0"))
            tables["trial_balance"].append({
                "account": str(row[0]), "name": row[1], "kind": row[2],
                "opening": op, "opening_dec": op_dec, "debit": de, "debit_dec": de_dec,
                "credit": cr, "credit_dec": cr_dec,
                "closing": float(Decimal(closing_dec)), "closing_dec": closing_dec,
                "closing_computed": True,  # file stores formula =D+E-F without cached value
                "row_id": r_no, "source_id": sid})
            parsed += 1
        # second sheet: Überleitung -> reconciliation rows
        if "Überleitung JA 2025" in wb.sheetnames:
            ws2 = wb["Überleitung JA 2025"]
            for r_no, row in enumerate(ws2.iter_rows(values_only=True), start=1):
                label = next((v for v in row if isinstance(v, str) and v.strip()), None)
                if not label:
                    continue
                value = next((v for v in row if isinstance(v, (int, float))), None)
                sid = _xl_register(registry, rel, fh, ws2.title, r_no, row)
                val, raw_a, canon = excel_amount(value)
                tables["reconciliation"].append({
                    "label": label, "value": val, "value_raw": raw_a, "value_dec": canon,
                    "file": rel, "row_id": r_no, "source_id": sid})
                parsed += 1
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=parsed, parsed_units=parsed,
                     parser_errors=["closing column has no cached values; computed as opening+debit-credit per the stored formula =D+E-F"],
                     parse_coverage="complete", population_scope=tb_scope)

    # --- Saldenliste 2024 (prior year) — structurally empty
    rel = f"{BEGLEIT}/Saldenliste_2024_Vorjahr.xlsx"
    path = data_dir / rel
    if not path.exists():
        _xlsx_absent(manifest, rel, "prior-year trial balance")
    else:
        fh = sha256_file(path)
        ws = openpyxl.load_workbook(path, data_only=True).active
        n_data = sum(1 for row in ws.iter_rows(min_row=4, values_only=True) if any(v is not None for v in row))
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=n_data, parsed_units=n_data, parser_errors=[],
                     population_scope="prior-year trial balance — EMPTY (0 data rows): no comparative-year data provided")

    # --- Abstimmung Nebenbücher/HB -> reconciliation
    rel = f"{BEGLEIT}/Abstimmung_Nebenbuecher_HB_2025.xlsx"
    path = data_dir / rel
    recon_scope = "management-prepared subledger/GL reconciliation per 31.12.2025 (zero differences claimed)"
    if not path.exists():
        _xlsx_absent(manifest, rel, recon_scope)
    else:
        fh = sha256_file(path)
        ws = openpyxl.load_workbook(path, data_only=True).active
        parsed = 0
        for r_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
            label = next((v for v in row if isinstance(v, str) and v.strip()), None)
            if not label:
                continue
            value = next((v for v in row if isinstance(v, (int, float))), None)
            sid = _xl_register(registry, rel, fh, ws.title, r_no, row)
            val, raw_a, canon = excel_amount(value)
            tables["reconciliation"].append({
                "label": label, "value": val, "value_raw": raw_a, "value_dec": canon,
                "file": rel, "row_id": r_no, "source_id": sid})
            parsed += 1
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=parsed, parsed_units=parsed, parser_errors=[],
                     population_scope=recon_scope)

    return tables, profile_feed


# ------------------------------------------------------------- docx/pdf units

def parse_doc_units(data_dir: Path, registry, manifest):
    doc_units = []

    # docx: paragraphs + table rows
    rel = f"{BEGLEIT}/Pruefungsplanung_JET_2025.docx"
    path = data_dir / rel
    # ROBUSTNESS (S1 missing-audit-plan): a missing audit-plan docx must not crash ingest.
    # It degrades to a 'failed' manifest entry with zero doc_units; the threshold extractor
    # then finds no doc-sourced values and falls back to config defaults (thresholds.py).
    if not path.exists():
        manifest.add(file=rel, file_hash=None, size_bytes=None, expected_units=None,
                     parsed_units=0, parser_errors=[f"{rel}: file not provided in this dossier"],
                     parse_coverage="failed", provided=False,
                     population_scope="normative source: JET audit planning — NOT PROVIDED "
                     "(thresholds fall back to config defaults)")
        return _parse_pdf_units(data_dir, registry, manifest, doc_units)
    fh = sha256_file(path)
    d = docx_lib.Document(path)
    parsed = 0
    for p_no, para in enumerate(d.paragraphs, start=1):
        if not para.text.strip():
            continue
        sid = registry.register(
            file=rel, file_hash=fh, kind="para",
            locator=f"{rel}#para:{p_no}", content=para.text,
            display_locator=f"{rel}:para {p_no}", para_no=p_no)
        doc_units.append({"source_id": sid, "file": rel, "page_index": None,
                          "page_label": None, "para_no": p_no,
                          "unit_type": "para", "text": para.text})
        parsed += 1
    for t_no, tbl in enumerate(d.tables, start=1):
        for r_no, trow in enumerate(tbl.rows, start=1):
            text = " | ".join(c.text for c in trow.cells)
            sid = registry.register(
                file=rel, file_hash=fh, kind="table",
                locator=f"{rel}#table:{t_no}#row:{r_no}", content=text,
                display_locator=f"{rel}:table {t_no}:row {r_no}", row_no=r_no)
            doc_units.append({"source_id": sid, "file": rel, "page_index": None,
                              "page_label": None, "para_no": None,
                              "unit_type": "table_row", "text": text})
            parsed += 1
    manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                 expected_units=parsed, parsed_units=parsed, parser_errors=[],
                 population_scope="normative source: JET audit planning (materiality 400k/300k, JET sampling threshold 25k, payment dual-approval limit 10k, selection criteria K1-K7)")

    return _parse_pdf_units(data_dir, registry, manifest, doc_units)


def _parse_pdf_units(data_dir: Path, registry, manifest, doc_units):
    """PDF text extraction (one 'page' registry unit per page + one 'para' doc_unit per
    text block). ROBUSTNESS (S1 and finals): a missing PDF degrades to a 'failed' manifest
    entry rather than crashing ingest."""
    pdf_scopes = {
        "Exportprotokoll_GDPdU_2025.pdf": "export protocol: declared row counts, amount sums and SHA-256 for the 8 GDPdU txt files",
        "IT-Bestaetigung_Vollstaendigkeit_2025.pdf": "IT completeness/immutability attestation (Soll=Haben=345,350,778.06; Festschreibung 2026-01-20)",
        "JA-Entwurf_2025_Auszug_Bilanz_GuV.pdf": "DRAFT financial statements extract (balance sheet + P&L) — audit object, unaudited",
    }
    for fname, scope in pdf_scopes.items():
        rel = f"{BEGLEIT}/{fname}"
        path = data_dir / rel
        if not path.exists():
            manifest.add(file=rel, file_hash=None, size_bytes=None, expected_units=None,
                         parsed_units=0, parser_errors=[f"{rel}: file not provided in this dossier"],
                         parse_coverage="failed", population_scope=scope + " — NOT PROVIDED",
                         provided=False)
            continue
        fh = sha256_file(path)
        pdf = fitz.open(path)
        parsed = 0
        for p_idx in range(len(pdf)):
            page = pdf[p_idx]
            label = str(p_idx + 1)
            registry.register(
                file=rel, file_hash=fh, kind="page",
                locator=f"{rel}#page:{p_idx}", content=page.get_text(),
                display_locator=f"{rel}:page {label}",
                page_index=p_idx, page_label=label)
            parsed += 1
            blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
            for b_no, b in enumerate(blocks, start=1):
                text = b[4].strip()
                sid = registry.register(
                    file=rel, file_hash=fh, kind="para",
                    locator=f"{rel}#page:{p_idx}#block:{b_no}", content=text,
                    display_locator=f"{rel}:page {label}:para {b_no}",
                    page_index=p_idx, page_label=label, para_no=b_no)
                doc_units.append({"source_id": sid, "file": rel, "page_index": p_idx,
                                  "page_label": label, "para_no": b_no,
                                  "unit_type": "para", "text": text})
        manifest.add(file=rel, file_hash=fh, size_bytes=path.stat().st_size,
                     expected_units=len(pdf), parsed_units=parsed, parser_errors=[],
                     population_scope=scope)
    return doc_units


def parse_sidecars(data_dir: Path, registry: SourceRegistry, manifest: ManifestBuilder):
    """Parse all 19 Begleitdokumente. Returns (tables, doc_units, csv_profile_feed)."""
    tables = {}
    profile_feed = []
    for table, spec in CSV_SPECS.items():
        rows, feed = _parse_csv_table(data_dir, table, spec, registry, manifest)
        tables[table] = rows
        profile_feed.append((feed[0], table, feed[1], feed[2]))
    xlsx_tables, xlsx_feed = parse_xlsx_files(data_dir, registry, manifest)
    tables.update(xlsx_tables)
    profile_feed.extend(xlsx_feed)
    doc_units = parse_doc_units(data_dir, registry, manifest)
    # Steuercodes: empty directory in the export
    st_dir = data_dir / "Steuercodes"
    if st_dir.exists() and not any(st_dir.iterdir()):
        manifest.add(file="Steuercodes/", file_hash=None, size_bytes=0,
                     expected_units=0, parsed_units=0, parser_errors=[],
                     population_scope="empty directory in the export — no tax-code tables provided")
    _list_unrecognized_begleit(data_dir, manifest)
    return tables, doc_units, profile_feed


# Files this ingest knows how to parse (used only to surface UNRECOGNIZED extras).
KNOWN_BEGLEIT_FILES = (
    {fname for fname, _ in CSV_SPECS.values()}
    | {"Berechtigungsauswertung_2025.xlsx", "OP-Liste_Debitoren_2025.xlsx",
       "OP-Liste_Kreditoren_2025.xlsx", "Saldenliste_2025.xlsx",
       "Saldenliste_2024_Vorjahr.xlsx", "Abstimmung_Nebenbuecher_HB_2025.xlsx"}
    | {"Pruefungsplanung_JET_2025.docx"}
    | {"Exportprotokoll_GDPdU_2025.pdf", "IT-Bestaetigung_Vollstaendigkeit_2025.pdf",
       "JA-Entwurf_2025_Auszug_Bilanz_GuV.pdf"}
)


def _list_unrecognized_begleit(data_dir: Path, manifest: ManifestBuilder):
    """ROBUSTNESS (S5 extra unknown file / S4 renamed file): surface any Begleitdokument
    that no adapter maps to as an explicit 'failed'-coverage manifest entry, so a finals
    file that was renamed or added is VISIBLE (coverage gap) rather than silently ignored —
    which would otherwise risk a phantom absence claim. No table/citation is produced."""
    beg = data_dir / BEGLEIT
    if not beg.is_dir():
        return
    for p in sorted(beg.iterdir()):
        if not p.is_file() or p.name.startswith(".") or p.name in KNOWN_BEGLEIT_FILES:
            continue
        manifest.add(file=f"{BEGLEIT}/{p.name}", file_hash=sha256_file(p),
                     size_bytes=p.stat().st_size, expected_units=None, parsed_units=0,
                     parser_errors=[f"{p.name}: unrecognized file — no adapter maps to it "
                                    "(present but UNPARSED; listed for coverage visibility)"],
                     parse_coverage="failed",
                     population_scope="UNRECOGNIZED file — not mapped to any parser; no "
                     "absence/existence claim may rely on its (non-)contents")
